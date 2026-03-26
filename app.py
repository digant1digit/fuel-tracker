import streamlit as st
import pandas as pd
from datetime import datetime
from io import BytesIO

# Excel
from openpyxl import Workbook
from openpyxl.chart import LineChart, Reference

# PDF
from reportlab.platypus import SimpleDocTemplate, Table, TableStyle, Paragraph
from reportlab.lib import colors
from reportlab.lib.styles import getSampleStyleSheet

# ------------------ PAGE CONFIG ------------------
st.set_page_config(page_title="Fuel Tracker", layout="wide")

st.title("⛽ Fuel Tracker")

# ------------------ SESSION INIT ------------------
if "df" not in st.session_state:
    st.session_state.df = pd.DataFrame(columns=["Date", "KM", "Liters", "Total Bill"])

# ------------------ TEMPLATE ------------------
sample_df = pd.DataFrame([
    {"Date": "2026-03-01", "KM": 1000, "Liters": 3, "Total Bill": 300}
])

st.download_button("⬇️ Download Template",
                   sample_df.to_csv(index=False).encode("utf-8"),
                   "fuel_template.csv")

st.info("📌 Supports: 01-03-2026 OR 2026-03-01")
st.warning("⚠️ Data is temporary. Download CSV to save your work.")

# ------------------ FILE UPLOAD ------------------
uploaded_file = st.file_uploader("📂 Upload CSV", type=["csv"])

if uploaded_file:
    df_uploaded = pd.read_csv(uploaded_file, sep=None, engine="python")
    st.session_state.df = df_uploaded
    st.success("CSV Loaded!")

df = st.session_state.df

# ------------------ VALIDATION ------------------
required = ["Date", "KM", "Liters", "Total Bill"]
if len(df) > 0 and not all(col in df.columns for col in required):
    st.error("❌ Invalid CSV format")
    st.stop()

# ------------------ CLEAN ------------------
if len(df) > 0:
    df["Date"] = pd.to_datetime(df["Date"], dayfirst=True, errors="coerce")

    for col in ["KM", "Liters", "Total Bill"]:
        df[col] = pd.to_numeric(df[col], errors="coerce")

    bad = df[df.isnull().any(axis=1)]
    if len(bad) > 0:
        st.warning("⚠️ Invalid rows removed")
        st.dataframe(bad, width="stretch")

    df = df.dropna()
    st.session_state.df = df

# ------------------ ADD ENTRY ------------------
st.subheader("➕ Add Entry")

with st.form("add"):
    c1, c2, c3, c4 = st.columns(4)

    date = c1.date_input("Date", datetime.today())
    km = c2.number_input("KM", min_value=0.0)
    liters = c3.number_input("Liters", min_value=0.0)
    amount = c4.number_input("₹", min_value=0.0)

    if st.form_submit_button("Add"):
        new_row = pd.DataFrame([{
            "Date": str(date),
            "KM": km,
            "Liters": liters,
            "Total Bill": amount
        }])

        st.session_state.df = pd.concat([st.session_state.df, new_row], ignore_index=True)
        st.success("Added!")
        st.rerun()

# ------------------ CALCULATIONS ------------------
df = st.session_state.df

if len(df) > 0:
    df = df.sort_values("Date")

    df["Distance"] = df["KM"].diff()
    df["Avg"] = df["Distance"] / df["Liters"]
    df["₹/KM"] = df["Total Bill"] / df["Distance"]
    df["Entry"] = range(1, len(df) + 1)

    total_d = df["Distance"].sum()
    total_l = df["Liters"].sum()
    overall = total_d / total_l if total_l else 0

    col1, col2 = st.columns(2)
    col1.metric("🚗 Last Avg", f"{df.iloc[-1]['Avg']:.2f} km/l")
    col2.metric("📊 Overall Avg", f"{overall:.2f} km/l")

    # ------------------ FILTER ------------------
    st.subheader("📅 Filter")

    c1, c2 = st.columns(2)
    start = c1.date_input("From", df["Date"].min())
    end = c2.date_input("To", df["Date"].max())

    df_f = df[(df["Date"] >= pd.to_datetime(start)) &
              (df["Date"] <= pd.to_datetime(end))]

    # ------------------ TABLE ------------------
    st.subheader("📋 Entries")

    show = df_f[["Entry", "Date", "KM", "Liters", "Total Bill", "Avg", "₹/KM"]]
    st.dataframe(show, width="stretch")

    # ------------------ EDIT / DELETE ------------------
    st.subheader("⚙️ Manage Entries")

    idx = st.selectbox("Select Entry", df.index)
    row = df.loc[idx]

    new_date = st.date_input("Edit Date", row["Date"])
    new_km = st.number_input("Edit KM", value=float(row["KM"]))
    new_liters = st.number_input("Edit Liters", value=float(row["Liters"]))
    new_amt = st.number_input("Edit ₹", value=float(row["Total Bill"]))

    col1, col2 = st.columns(2)

    if col1.button("Update"):
        st.session_state.df.loc[idx, ["Date", "KM", "Liters", "Total Bill"]] = [
            str(new_date), new_km, new_liters, new_amt
        ]
        st.success("Updated")
        st.rerun()

    if col2.button("Delete"):
        st.session_state.df = st.session_state.df.drop(idx).reset_index(drop=True)
        st.warning("Deleted")
        st.rerun()

    # ------------------ GRAPH ------------------
    st.subheader("📈 Monthly Trend")

    df["Month"] = df["Date"].dt.to_period("M").astype(str)
    monthly = df.groupby("Month")["Avg"].mean()

    st.line_chart(monthly)

    # ------------------ EXPORT PREP ------------------
    export = df.copy()
    export["Avg"] = export["Avg"].fillna(0).round(2)
    export["₹/KM"] = export["₹/KM"].fillna(0).round(2)

    # ------------------ CSV EXPORT ------------------
    st.subheader("💾 Export")

    st.download_button(
        "📥 Download CSV",
        export[["Date","KM","Liters","Total Bill","Avg","₹/KM"]].to_csv(index=False),
        "fuel_data.csv"
    )

    # ------------------ EXCEL EXPORT ------------------
    wb = Workbook()
    ws = wb.active

    headers = ["Date","KM","Liters","Total Bill","Avg","₹/KM"]
    ws.append(headers)

    for _, r in export.iterrows():
        ws.append([str(r["Date"]), r["KM"], r["Liters"], r["Total Bill"], r["Avg"], r["₹/KM"]])

    chart = LineChart()
    data = Reference(ws, min_col=5, min_row=1, max_row=ws.max_row)
    chart.add_data(data, titles_from_data=True)
    ws.add_chart(chart, "H2")

    buf = BytesIO()
    wb.save(buf)

    st.download_button("📊 Download Excel", buf.getvalue(), "fuel_report.xlsx")

    # ------------------ PDF EXPORT ------------------
    pdf_buf = BytesIO()
    doc = SimpleDocTemplate(pdf_buf)
    styles = getSampleStyleSheet()

    elements = [Paragraph("Fuel Report", styles["Title"])]

    table_data = [["Date","KM","Liters","Bill","Avg","₹/KM"]]
    for _, r in export.iterrows():
        table_data.append([str(r["Date"]), r["KM"], r["Liters"], r["Total Bill"], r["Avg"], r["₹/KM"]])

    table = Table(table_data)
    table.setStyle(TableStyle([
        ("BACKGROUND",(0,0),(-1,0),colors.grey),
        ("TEXTCOLOR",(0,0),(-1,0),colors.white),
        ("GRID",(0,0),(-1,-1),1,colors.black)
    ]))

    elements.append(table)
    doc.build(elements)

    st.download_button("📄 Download PDF", pdf_buf.getvalue(), "fuel_report.pdf")
